"""
pipeline_runner.py — รัน ERP pipeline และคืน path ไฟล์ผลลัพธ์
"""
import subprocess
import sys
import os
import shutil
import tempfile
from pathlib import Path

SCRIPTS_DIR = Path(__file__).parent / "scripts"


def _run(cmd: list, log_lines: list):
    """รันคำสั่งและเก็บ log"""
    result = subprocess.run(
        cmd, capture_output=True, text=True
    )
    stdout = result.stdout.strip()
    stderr = result.stderr.strip()
    if stdout:
        log_lines.append(stdout)
    if stderr:
        log_lines.append(f"[WARN] {stderr}")
    if result.returncode != 0:
        raise RuntimeError(
            f"Script ล้มเหลว: {' '.join(cmd)}\n{stderr}"
        )


def run_pipeline(
    erp_path: str,
    pps_path: str | None,
    password: str,
    month: int,
    year: int,
    progress_cb=None,
) -> tuple[str, list[str]]:
    """
    รัน ERP pipeline ทั้งหมด
    คืนค่า (output_path, log_lines)
    """
    log = []
    tmp = tempfile.mkdtemp(prefix="erp_")
    python = sys.executable

    def step(msg, fn, *args):
        if progress_cb:
            progress_cb(msg)
        log.append(f"▶ {msg}")
        fn(*args)

    try:
        decrypted = os.path.join(tmp, "decrypted.xlsx")
        step("ถอดรหัสไฟล์ ERP...", _run,
             [python, str(SCRIPTS_DIR / "decrypt_erp.py"),
              erp_path, decrypted, password], log)

        # ตรวจรูปแบบไฟล์
        import openpyxl
        wb = openpyxl.load_workbook(decrypted, read_only=True)
        ws2 = wb["2"] if "2" in wb.sheetnames else None
        is_new_format = False
        if ws2:
            first_cell = ws2.cell(1, 1).value
            is_new_format = first_cell is not None and str(first_cell).strip() == ""
            # ตรวจอีกแบบ: แถวที่ 1 มีหลายคอลัมน์ว่าง = format ใหม่
            row1_vals = [ws2.cell(1, c).value for c in range(1, 6)]
            if all(v is None for v in row1_vals):
                is_new_format = True
        wb.close()

        current = decrypted

        if is_new_format:
            preprocessed = os.path.join(tmp, "preprocessed.xlsx")
            step("ปรับ format ERP ใหม่...", _run,
                 [python, str(SCRIPTS_DIR / "preprocess_erp.py"),
                  current, preprocessed], log)
            current = preprocessed

        if pps_path:
            base = os.path.join(tmp, "base.xlsx")
            step("สร้างชีท PPS จากไฟล์แยก...", _run,
                 [python, str(SCRIPTS_DIR / "build_pps.py"),
                  pps_path, current, base], log)
            current = base

        with_t = os.path.join(tmp, "with_t.xlsx")
        step("สร้างชีท T2-T18...", _run,
             [python, str(SCRIPTS_DIR / "build_t_sheets.py"),
              current, with_t], log)
        current = with_t

        # ถ้าไม่มี PPS ภายนอก ให้ build จาก ERP
        import openpyxl as ox
        wb2 = ox.load_workbook(current, read_only=True)
        has_pps = any(s.lower() == "pps" for s in wb2.sheetnames)
        has_pps05 = any(s.lower() == "pps05" for s in wb2.sheetnames)
        wb2.close()

        if not pps_path and has_pps and not has_pps05:
            base_pps = os.path.join(tmp, "base_pps.xlsx")
            step("สร้างชีท pps05 จาก ERP...", _run,
                 [python, str(SCRIPTS_DIR / "build_pps_from_erp.py"),
                  current, base_pps], log)
            # ลบชีท PPS ดิบออก
            wb3 = ox.load_workbook(base_pps)
            if "PPS" in wb3.sheetnames:
                del wb3["PPS"]
            clean_path = os.path.join(tmp, "base_pps_clean.xlsx")
            wb3.save(clean_path)
            current = clean_path

        with_cal = os.path.join(tmp, "with_cal.xlsx")
        step("คำนวณ CAL (ค่าคอม)...", _run,
             [python, str(SCRIPTS_DIR / "build_cal05.py"),
              current, with_cal], log)
        current = with_cal

        tmp_sum = os.path.join(tmp, "tmp_sum.xlsx")
        step("สร้างชีท sum...", _run,
             [python, str(SCRIPTS_DIR / "build_sum.py"),
              current, tmp_sum, "sum"], log)

        output_name = f"รายงานขายค่าคอมเดือน{month}_{year}.xlsx"
        output_path = os.path.join(tmp, output_name)
        step("สร้างชีท sum (ตัดO2O)...", _run,
             [python, str(SCRIPTS_DIR / "build_sum.py"),
              tmp_sum, output_path, "sum (ตัดO2O)"], log)

        # Recalculate ด้วย LibreOffice
        step("Recalculate สูตรด้วย LibreOffice...", _run,
             ["soffice", "--headless", "--convert-to", "xlsx",
              "--outdir", tmp, output_path], log)

        log.append(f"✅ เสร็จสิ้น: {output_name}")
        return output_path, log

    except Exception as e:
        log.append(f"❌ ERROR: {e}")
        raise

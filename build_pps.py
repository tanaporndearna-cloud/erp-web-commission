"""
อ่านไฟล์ PPS.xlsx (export จากระบบ) แล้ว inject เป็นชีต PPS มาตรฐานใน ERP xlsx

โครงสร้าง output ชีต PPS:
  Col A  Product Code       ← source col C
  Col B  Description        ← source col D
  Col C  Category           ← source col B
  Col D  (ว่าง)
  Col E  Branch Province    ← source col O (15)
  Col F  Retail Price       ← source col P (16)
  Col G  Wholesale General  ← source col T (20)
  Col H  P Retail Price     = =IF(E{r}>0,E{r},F{r})
  Col I  Wholesale General P = =(G{r}+H{r})/2

ชีต PPS เดิมในไฟล์ ERP (ถ้ามี) จะถูก replace ด้วยชีตนี้
ชีตใหม่จะชื่อ "PPS" เสมอ (ไม่ขึ้นกับเดือน)

Usage:
    python3 build_pps.py <pps_source.xlsx> <erp_working.xlsx> <output.xlsx>
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill

PPS_PATH = sys.argv[1]   # PPS.xlsx จากระบบ
ERP_PATH = sys.argv[2]   # ไฟล์ ERP ที่ผ่าน decrypt แล้ว
OUT_PATH = sys.argv[3]   # output

# --- อ่านข้อมูลจาก PPS.xlsx ---
print(f'Reading PPS source: {PPS_PATH}')
pps_wb = openpyxl.load_workbook(PPS_PATH, data_only=True)
pps_ws = pps_wb.active
pps_last = pps_ws.max_row
print(f'  {pps_last} rows, {pps_ws.max_column} cols')

# source column index (1-based)
SRC_PRODUCT_CODE    = 3   # C
SRC_DESCRIPTION     = 4   # D
SRC_CATEGORY        = 2   # B
SRC_BRANCH_PROVINCE = 15  # O
SRC_RETAIL_PRICE    = 16  # P
SRC_WHOLESALE_GEN   = 20  # T

# --- โหลดไฟล์ ERP ---
print(f'Loading ERP file: {ERP_PATH}')
wb = openpyxl.load_workbook(ERP_PATH)

# ลบชีต PPS เดิม (ถ้ามี)
for name in wb.sheetnames:
    if 'pps' in name.lower():
        del wb[name]
        print(f'  Removed old sheet: "{name}"')

# สร้างชีต PPS ใหม่
ws = wb.create_sheet('PPS')
print(f'  Created new sheet "PPS"')

# header row
HEADERS = ['Product Code', 'Description', 'Category', '',
           'Branch Province', 'Retail Price', 'Wholesale General',
           'P Retail Price', 'Wholesale General P']
for c, h in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=c, value=h if h else None)
    if h:
        cell.font = Font(bold=True)

# เขียนข้อมูล (ข้ามแถว header ของ source = row 1)
out_row = 2
for r in range(2, pps_last + 1):
    # ข้ามแถวที่ Product Code ว่าง
    if pps_ws.cell(row=r, column=SRC_PRODUCT_CODE).value is None:
        continue

    ws.cell(row=out_row, column=1, value=pps_ws.cell(row=r, column=SRC_PRODUCT_CODE).value)
    ws.cell(row=out_row, column=2, value=pps_ws.cell(row=r, column=SRC_DESCRIPTION).value)
    ws.cell(row=out_row, column=3, value=pps_ws.cell(row=r, column=SRC_CATEGORY).value)
    # col D ว่าง
    ws.cell(row=out_row, column=5, value=pps_ws.cell(row=r, column=SRC_BRANCH_PROVINCE).value)
    ws.cell(row=out_row, column=6, value=pps_ws.cell(row=r, column=SRC_RETAIL_PRICE).value)
    ws.cell(row=out_row, column=7, value=pps_ws.cell(row=r, column=SRC_WHOLESALE_GEN).value)
    ws.cell(row=out_row, column=8, value=f'=IF(E{out_row}>0,E{out_row},F{out_row})')
    ws.cell(row=out_row, column=9, value=f'=(G{out_row}+H{out_row})/2')
    out_row += 1

data_rows = out_row - 2
print(f'  Written {data_rows} data rows to "PPS"')

wb.save(OUT_PATH)
print(f'saved {OUT_PATH}')

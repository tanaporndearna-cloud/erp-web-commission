"""
Preprocess raw ERP export (format ใหม่ จาก eFlowsys) สำหรับชีต 2-18:

1. ลบแถว 1-12 ออก (logo / header)
2. เอา merge cells และ freeze panes ออก
3. ลบแถวสุดท้ายที่มีคำว่า "รวมทั้งหมด" ออก
4. ลบคอลัมน์ที่ว่างทั้งหมดออก
5. แทรกคอลัมน์แรก ชื่อ "สาขา" ใส่ค่าตามชีต (ชีต 2=T2, ชีต 18=T18)
6. ลบรูปภาพ/โลโก้ทั้งหมดในไฟล์ออก
7. ปรับความกว้างคอลัมน์ให้เหมาะสมกับข้อมูล และปรับความสูงทุกแถวเป็น 15pt
8. เปิด wrap text ให้คอลัมน์ ชื่อลูกค้า / รายละเอียด / สาขา

Usage:
    python3 preprocess_erp.py <input.xlsx> <output.xlsx>
"""
import sys
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

IN_PATH = sys.argv[1]
OUT_PATH = sys.argv[2]

wb = openpyxl.load_workbook(IN_PATH)

sheet_names = [s for s in wb.sheetnames if s.isdigit() and 2 <= int(s) <= 18]

for sname in sheet_names:
    ws = wb[sname]
    n = int(sname)
    branch_name = f'T{n}'
    print(f"Processing sheet {sname} ...")

    # Step 6: ลบรูปภาพ/โลโก้
    ws._images = []

    # Step 2: Unmerge ทุก merged cells
    for mc in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(mc))

    # Step 2: ลบ freeze panes
    ws.freeze_panes = None

    # Step 1: ลบแถว 1-12
    ws.delete_rows(1, 12)
    # ตอนนี้ row 1 = เดิม row 13 (header), data เริ่ม row 2

    # Step 3: ลบแถว "รวมทั้งหมด"
    for r in range(ws.max_row, 0, -1):
        v = ws.cell(row=r, column=1).value
        if v and isinstance(v, str) and 'รวม' in v:
            ws.delete_rows(r)
            break

    # Step 4: ลบคอลัมน์ที่ว่าง (ตรวจที่แถวข้อมูล row 2+)
    cols_to_delete = []
    for c in range(1, ws.max_column + 1):
        has_data = False
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=c).value is not None:
                has_data = True
                break
        if not has_data:
            cols_to_delete.append(c)
    for c in reversed(cols_to_delete):
        ws.delete_cols(c)

    # Step 5: แทรกคอลัมน์ "สาขา" ที่ตำแหน่ง 1
    ws.insert_cols(1)
    ws.cell(row=1, column=1, value='สาขา')
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=2).value is not None:
            ws.cell(row=r, column=1, value=branch_name)

    # Step 7: ความสูงแถว 15pt
    for r in range(1, ws.max_row + 1):
        ws.row_dimensions[r].height = 15

    # Step 7 & 8: ความกว้างคอลัมน์ + wrap text
    WRAP_HEADERS = {'ชื่อลูกค้า', 'รายละเอียด', 'สาขา'}
    for c in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=c).value
        col_letter = get_column_letter(c)
        should_wrap = (header in WRAP_HEADERS) if header else False

        max_len = 0
        for r in range(1, ws.max_row + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                cell_len = len(str(v))
                if cell_len > max_len:
                    max_len = cell_len

        ws.column_dimensions[col_letter].width = max(8, min(max_len + 2, 40))

        if should_wrap:
            for r in range(1, ws.max_row + 1):
                cell = ws.cell(row=r, column=c)
                old = cell.alignment
                cell.alignment = Alignment(
                    wrap_text=True,
                    horizontal=old.horizontal,
                    vertical=old.vertical,
                )

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    print(f"  -> {ws.max_row - 1} data rows, {ws.max_column} cols | headers: {headers}")

wb.save(OUT_PATH)
print(f"\nSaved to {OUT_PATH}")

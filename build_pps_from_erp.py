"""
Convert raw PPS sheet (already inside ERP file) to standard format (pps05).

Raw PPS columns:
  Col 2  = Category
  Col 3  = Product Code (lookup key)
  Col 4  = Description
  Col 16 = Retail Price        -> F (6)
  Col 20 = Wholesale General   -> G (7)
  Col 17 = ราคาพิเศษ           -> E (5) เก็บค่า numeric เท่านั้น

Standard pps05 format (expected by build_cal05.py):
  A (1) = Product Code
  B (2) = Description
  C (3) = Category
  D (4) = (empty)
  E (5) = ราคาพิเศษ (numeric only; blank = ไม่มีราคาพิเศษ)
  F (6) = Retail Price
  G (7) = Wholesale General
  H (8) = P Retail Price  — สูตร =IF(E>0,E,F)  (ราคาพิเศษถ้ามี มิฉะนั้น Retail)
  I (9) = P Wholesale General — สูตร =(G+H)/2
"""
import sys
import openpyxl

IN_PATH = sys.argv[1]
OUT_PATH = sys.argv[2]

wb = openpyxl.load_workbook(IN_PATH)

raw_ws = wb['PPS']

# Remove existing pps05 if any
for name in ['pps05', 'PPS05', 'pps']:
    if name in wb.sheetnames and name != 'PPS':
        del wb[name]

# Create pps05 sheet
pps_ws = wb.create_sheet('pps05')

# Header
pps_ws.cell(row=1, column=1, value='Product Code')
pps_ws.cell(row=1, column=2, value='Description')
pps_ws.cell(row=1, column=3, value='Category')
pps_ws.cell(row=1, column=4, value='')
pps_ws.cell(row=1, column=5, value='ราคาพิเศษ')
pps_ws.cell(row=1, column=6, value='Retail Price')
pps_ws.cell(row=1, column=7, value='Wholesale General')
pps_ws.cell(row=1, column=8, value='P Retail Price')
pps_ws.cell(row=1, column=9, value='P Wholesale General')

out_row = 2
for r in range(2, raw_ws.max_row + 1):
    prod_code = raw_ws.cell(row=r, column=3).value
    if not prod_code:
        continue
    category = raw_ws.cell(row=r, column=2).value
    desc     = raw_ws.cell(row=r, column=4).value
    retail   = raw_ws.cell(row=r, column=16).value
    wholesale= raw_ws.cell(row=r, column=20).value
    p_retail = raw_ws.cell(row=r, column=17).value  # ราคาพิเศษ (col 17)

    # E = ราคาพิเศษ: เก็บเฉพาะตัวเลข (ข้อความ เช่น "PROA" ให้ใส่ None)
    p_retail_num = p_retail if isinstance(p_retail, (int, float)) else None

    pps_ws.cell(row=out_row, column=1, value=prod_code)
    pps_ws.cell(row=out_row, column=2, value=desc)
    pps_ws.cell(row=out_row, column=3, value=category)
    pps_ws.cell(row=out_row, column=4, value=None)
    pps_ws.cell(row=out_row, column=5, value=p_retail_num)   # E = ราคาพิเศษ
    pps_ws.cell(row=out_row, column=6, value=retail)          # F = Retail Price
    pps_ws.cell(row=out_row, column=7, value=wholesale)       # G = Wholesale General
    # H = =IF(E>0,E,F)  ราคาพิเศษถ้ามี มิฉะนั้น Retail Price
    pps_ws.cell(row=out_row, column=8,
                value=f'=IF(E{out_row}>0,E{out_row},F{out_row})')
    # I = =(G+H)/2  ค่าเฉลี่ย Wholesale และ P Retail
    pps_ws.cell(row=out_row, column=9,
                value=f'=(G{out_row}+H{out_row})/2')
    out_row += 1

print(f"pps05 sheet created with {out_row-2} rows")

wb.save(OUT_PATH)
print(f"Saved to {OUT_PATH}")
